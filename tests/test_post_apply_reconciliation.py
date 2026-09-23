"""Post-apply reconciliation: native md/txt bodies round-trip through
ingestion, re-ingestion keeps the managed-note relationship, and a vault
note that already equals the new generated body is IN_SYNC (and only
re-baselined, never rewritten or reported as a conflict)."""
import asyncio
import io
import json
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from ingest.apply import ApplyRefused, build_candidate
from ingest.markdown import document_body, normalize, to_markdown
from ingest.reverse_analyzer import analyze_managed_notes
from ingest.sync import BODY_CONTRACT_VERSION, index_local_path
from ingest.vault_writer import hash_managed_body, parse_frontmatter, write_managed_note
from main import index_to_qdrant
from sync_cli import load_state
from tests.test_apply import HUMAN_BODY, MD_SOURCE, ApplyFixture, _quiet
from tests.test_main_cli import make_store
from tests.test_vault_writer import FakeObsidian, default_metadata

NATIVE_MD = ("---\ntitle: Native\nproject: demo\ntags: [alpha, beta]\n---\n\n"
             "# Heading\n\nA short paragraph.\n\n- existing item\n1. numbered\n\n"
             "```python\nprint('x')\n```\n\n> quote\n")
NATIVE_MD_BODY = NATIVE_MD.split("---\n\n", 1)[1]


def _minimal_pdf(text):
    """A tiny valid single-page PDF whose content stream draws `text`."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


class NativeBodyContractTests(unittest.TestCase):
    def setUp(self):
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)

    def ingest(self, name, content, state=None, store=None):
        path = self.root / name
        if isinstance(content, bytes):
            path.write_bytes(content)
        else:
            path.write_text(content, encoding="utf-8")
        report, state = index_local_path(store or make_store(), str(path), state={} if state is None else state)
        return path, report, state[str(path.resolve())]

    def test_markdown_structure_is_preserved_verbatim(self):
        _, _, entry = self.ingest("native.md", NATIVE_MD)
        self.assertEqual(entry["generated_body"], NATIVE_MD_BODY)
        self.assertEqual(entry["body_contract"], BODY_CONTRACT_VERSION)

    def test_headings_are_not_bulleted(self):
        body = document_body("# Title\n## Sub\nshort line\n", "md")
        self.assertEqual(body, "# Title\n## Sub\nshort line\n")
        self.assertNotIn("- #", body)

    def test_markdown_frontmatter_is_source_metadata_not_body(self):
        _, report, entry = self.ingest("native.md", NATIVE_MD)
        self.assertNotIn("- ---", entry["generated_body"])
        self.assertNotIn("title: Native", entry["generated_body"])
        # Exactly one frontmatter block (the ingestion one) in the indexed text.
        indexed = report["indexed"][0]["markdown"]
        self.assertEqual(indexed.count("---\n"), 2)
        # Source frontmatter now reaches filter metadata.
        self.assertEqual(entry["filter_metadata"]["project"], "demo")
        self.assertEqual(entry["filter_metadata"]["tags"], ["alpha", "beta"])

    def test_filter_metadata_is_stable_on_unchanged_skip(self):
        store = make_store()
        path, _, entry = self.ingest("native.md", NATIVE_MD, store=store)
        state = {str(path.resolve()): entry}
        report, _ = index_local_path(store, str(path), state=state)
        self.assertEqual(report["skipped_unchanged"], [str(path.resolve())])
        self.assertEqual(report["metadata_updated"], [])

    def test_txt_is_verbatim_and_deterministic(self):
        text = "Line one\n  indented\n\n---\nnot frontmatter\n"
        _, _, first = self.ingest("a.txt", text)
        _, _, second = self.ingest("b.txt", text)
        self.assertEqual(first["generated_body"], text)
        self.assertEqual(first["generated_body"], second["generated_body"])

    def test_bom_is_dropped_for_native_types(self):
        self.assertEqual(document_body("﻿---\na: b\n---\n\nbody\n", "md"), "body\n")
        self.assertEqual(document_body("﻿plain\n", "txt"), "plain\n")

    def test_legacy_md_entry_is_reprocessed_once_to_the_new_contract(self):
        store = make_store()
        path, _, entry = self.ingest("native.md", NATIVE_MD, store=store)
        legacy = dict(entry, generated_body=normalize(NATIVE_MD))
        del legacy["body_contract"]
        state = {str(path.resolve()): legacy}
        report, state = index_local_path(store, str(path), state=state)
        self.assertEqual(len(report["indexed"]), 1)
        self.assertEqual(state[str(path.resolve())]["generated_body"], NATIVE_MD_BODY)
        report, _ = index_local_path(store, str(path), state=state)
        self.assertEqual(len(report["skipped_unchanged"]), 1)

    def test_pdf_and_docx_bodies_are_unchanged(self):
        from docx import Document
        doc = Document()
        for line in ("DOCX heading", "second paragraph", "x" * 130):
            doc.add_paragraph(line)
        doc.save(self.root / "d.docx")
        docx_path, _, docx_entry = self.ingest("d.docx", (self.root / "d.docx").read_bytes())
        self.assertEqual(docx_entry["generated_body"],
                         normalize("DOCX heading\nsecond paragraph\n" + "x" * 130))
        self.assertTrue(docx_entry["generated_body"].startswith("- DOCX heading\n- second paragraph"))

        pdf_path, _, pdf_entry = self.ingest("p.pdf", _minimal_pdf("PDF body text"))
        from ingest.extract import extract_text
        self.assertEqual(pdf_entry["generated_body"], normalize(extract_text(str(pdf_path), "pdf")))
        self.assertIn("- PDF body text", pdf_entry["generated_body"])
        self.assertEqual(to_markdown("a\nb", "/x.pdf", "pdf").split("---\n\n", 1)[1], "- a\n- b")
        self.assertEqual(to_markdown("a\nb", "/x").split("---\n\n", 1)[1], "- a\n- b")  # no ftype: unchanged

    def test_pdf_and_docx_are_not_forced_to_reprocess_by_the_contract(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("unchanged docx")
        doc.save(self.root / "d.docx")
        store = make_store()
        path, _, entry = self.ingest("d.docx", (self.root / "d.docx").read_bytes(), store=store)
        legacy = dict(entry)
        del legacy["body_contract"]
        report, _ = index_local_path(store, str(path), state={str(path.resolve()): legacy})
        self.assertEqual(len(report["skipped_unchanged"]), 1)


class VaultWriteRelationshipTests(unittest.TestCase):
    def setUp(self):
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.state_path = self.root / "sync_state.json"
        self.source = self.root / "input" / "doc.md"
        self.source.parent.mkdir()
        self.source.write_text("# Doc\n\nfirst version\n")
        self.store = make_store()
        self.ob = FakeObsidian()
        _quiet(index_to_qdrant, str(self.source), None, self.state_path, store=self.store,
               write_vault=True, obsidian=self.ob)
        self.key = str(self.source.resolve())
        self.dest = load_state(self.state_path, namespace="local")[self.key]["vault_write"]["dest_path"]

    def reingest(self, **kwargs):
        return _quiet(index_to_qdrant, str(self.source), None, self.state_path, store=self.store, **kwargs)

    def entry(self):
        return load_state(self.state_path, namespace="local")[self.key]

    def test_changed_source_reingest_keeps_relationship_but_not_stale_status(self):
        before = self.entry()["vault_write"]
        self.assertEqual(before["status"], "created")
        self.source.write_text("# Doc\n\nsecond version\n")
        self.reingest()
        after = self.entry()["vault_write"]
        self.assertEqual(after, {"dest_path": self.dest, "status": "unverified", "generated_sha256": None})

    def test_analyzer_still_finds_note_after_reingest(self):
        self.source.write_text("# Doc\n\nsecond version\n")
        self.reingest()
        results = asyncio.run(analyze_managed_notes(self.ob, load_state(self.state_path, namespace="local")))
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["managed_note_path"], self.dest)
        # The note still holds the untouched first version: the source moved
        # on, nobody edited the vault. Never HUMAN_MODIFIED (which could be
        # applied and revert the source to the stale generated text).
        self.assertEqual(results[0]["classification"], "SOURCE_CHANGED")
        self.assertTrue(results[0]["flags"]["vault_generated_from_older_source"])

    def test_unedited_outdated_note_never_yields_an_applicable_proposal(self):
        from main import propose_vault_changes
        self.source.write_text("# Doc\n\nsecond version\n")
        self.reingest()
        buf = io.StringIO()
        with redirect_stdout(buf):
            propose_vault_changes(self.state_path, proposals_dir=self.root / "p", json_output=True, obsidian=self.ob)
        self.assertEqual([c["classification"] for c in json.loads(buf.getvalue())], ["SOURCE_CHANGED"])

    def test_human_edit_plus_source_change_is_human_modified_with_flag(self):
        content = self.ob.files[self.dest]
        self.ob.files[self.dest] = content[:content.index("---\n\n") + 5] + "a human edit\n"
        self.source.write_text("# Doc\n\nsecond version\n")
        self.reingest()
        [result] = asyncio.run(analyze_managed_notes(self.ob, load_state(self.state_path, namespace="local")))
        self.assertEqual(result["classification"], "HUMAN_MODIFIED")
        self.assertTrue(result["flags"]["vault_generated_from_older_source"])

    def test_unmodified_note_is_updated_normally_after_source_change(self):
        self.source.write_text("# Doc\n\nsecond version\n")
        self.reingest(write_vault=True, obsidian=self.ob)
        self.assertEqual(self.entry()["vault_write"]["status"], "updated")
        self.assertIn("second version", self.ob.files[self.dest])


class RebaselineUnitTests(unittest.IsolatedAsyncioTestCase):
    async def test_live_body_equal_to_generated_rebaselines_without_touching_body(self):
        ob = FakeObsidian()
        await write_managed_note(ob, "M/n.md", "old generated\n", default_metadata())
        fm_end = ob.files["M/n.md"].index("---\n\n") + 5
        ob.files["M/n.md"] = ob.files["M/n.md"][:fm_end] + "human text\n\n"
        result = await write_managed_note(ob, "M/n.md", "human text\n", default_metadata(source_sha256="new"))
        self.assertEqual(result["status"], "rebaselined")
        fm, body = parse_frontmatter(ob.files["M/n.md"])
        self.assertEqual(body, "human text\n\n")  # byte-for-byte, including outer whitespace
        self.assertEqual(fm["cortex_generated_sha256"], hash_managed_body("human text"))
        self.assertEqual(fm["cortex_source_sha256"], "new")
        # Now a genuine later change is a normal safe update, not a conflict.
        self.assertEqual((await write_managed_note(ob, "M/n.md", "next\n", default_metadata()))["status"],
                         "updated")

    async def test_non_cortex_frontmatter_is_never_rewritten(self):
        ob = FakeObsidian()
        await write_managed_note(ob, "M/n.md", "old generated\n", default_metadata())
        content = ob.files["M/n.md"].replace("cortex_managed: true", "cortex_managed: true\naliases: mine", 1)
        fm_end = content.index("---\n\n") + 5
        ob.files["M/n.md"] = content[:fm_end] + "human text\n"
        ob.write_calls.clear()
        result = await write_managed_note(ob, "M/n.md", "human text\n", default_metadata())
        self.assertEqual(result["status"], "skipped")
        self.assertEqual(ob.write_calls, [])
        self.assertIn("aliases: mine", ob.files["M/n.md"])

    async def test_genuine_human_edit_still_conflicts(self):
        ob = FakeObsidian()
        await write_managed_note(ob, "M/n.md", "generated\n", default_metadata())
        fm_end = ob.files["M/n.md"].index("---\n\n") + 5
        ob.files["M/n.md"] = ob.files["M/n.md"][:fm_end] + "human text\n"
        ob.write_calls.clear()
        result = await write_managed_note(ob, "M/n.md", "different generated\n", default_metadata())
        self.assertEqual(result["status"], "conflict")
        self.assertEqual(ob.write_calls, [])


class PostApplyLifecycleTests(ApplyFixture):
    """source -> ingest -> vault -> human edit -> propose -> approve -> apply
    -> normal re-ingestion -> consistent again."""

    def analyze(self):
        return asyncio.run(analyze_managed_notes(self.ob, load_state(self.state_path, namespace="local")))

    def reingest(self, **kwargs):
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = index_to_qdrant(str(self.source), None, self.state_path, store=self.store, **kwargs)
        return code, buf.getvalue()

    def entry(self):
        return load_state(self.state_path, namespace="local")[str(self.source.resolve())]

    def test_markdown_round_trip_converges_to_in_sync(self):
        self.seed()
        self.assertTrue(self.apply()["ok"])
        live_hash = self.proposal()["fingerprints"]["live_vault_sha256"]
        _, out = self.reingest()
        self.assertIn("Indexed 1 file(s)", out)
        entry = self.entry()
        self.assertEqual(entry["generated_body"], HUMAN_BODY)
        self.assertEqual(hash_managed_body(entry["generated_body"]), live_hash)
        self.assertEqual(entry["vault_write"]["dest_path"], self.note_path)
        [result] = self.analyze()
        self.assertEqual(result["classification"], "IN_SYNC", result)
        self.assertFalse(result["flags"]["source_changed"])

    def test_write_vault_after_round_trip_rebaselines_without_conflict_or_body_rewrite(self):
        self.seed()
        self.apply()
        self.reingest()
        _, body_before = parse_frontmatter(self.ob.files[self.note_path])
        code, out = self.reingest(write_vault=True, obsidian=self.ob)
        self.assertEqual(code, 0)
        self.assertNotIn("CONFLICT", out)
        self.assertIn("baseline recorded", out)
        fm, body_after = parse_frontmatter(self.ob.files[self.note_path])
        self.assertEqual(body_after, body_before)
        self.assertEqual(fm["cortex_generated_sha256"], hash_managed_body(HUMAN_BODY))
        self.assertEqual(self.entry()["vault_write"]["status"], "rebaselined")
        # A second run has nothing to do at all.
        self.ob.write_calls.clear()
        _, out = self.reingest(write_vault=True, obsidian=self.ob)
        self.assertEqual(self.ob.write_calls, [])
        self.assertIn("unchanged, skipped", out)
        self.assertEqual(self.analyze()[0]["classification"], "IN_SYNC")

    def test_new_human_edit_after_reconciliation_is_detected_and_protected(self):
        self.seed()
        self.apply()
        self.reingest(write_vault=True, obsidian=self.ob)
        self.human_edit("# Human heading\n\nA second, newer human edit.\n")
        self.assertEqual(self.analyze()[0]["classification"], "HUMAN_MODIFIED")
        _, out = self.reingest(write_vault=True, obsidian=self.ob)
        self.assertIn("CONFLICT", out)
        self.assertIn("newer human edit", self.ob.files[self.note_path])

    def test_later_source_change_after_reconciliation_updates_normally(self):
        self.seed()
        self.apply()
        self.reingest(write_vault=True, obsidian=self.ob)
        self.source.write_text(self.source.read_text() + "\nAppended in the source.\n")
        _, out = self.reingest(write_vault=True, obsidian=self.ob)
        self.assertIn("Vault: updated", out)
        self.assertIn("Appended in the source.", self.ob.files[self.note_path])

    def test_txt_round_trip_converges_to_in_sync(self):
        self.seed(name="notes.txt", content="plain original\n", human_body="Edited text\n\nsecond para\n")
        self.assertTrue(self.apply()["ok"])
        self.assertEqual(self.source.read_text(), "Edited text\n\nsecond para\n")
        self.reingest()
        self.assertEqual(self.entry()["generated_body"], "Edited text\n\nsecond para\n")
        self.assertEqual(self.analyze()[0]["classification"], "IN_SYNC")

    def test_markdown_source_frontmatter_survives_the_full_cycle(self):
        self.seed()
        self.apply()
        self.reingest(write_vault=True, obsidian=self.ob)
        self.assertTrue(self.source.read_text().startswith("---\ntitle: Keep Me\ntags: alpha\n---\n\n"))
        note = self.ob.files[self.note_path]
        self.assertNotIn("title: Keep Me", note)  # source metadata never enters the managed note
        self.assertEqual(note.count("---\n"), 2)   # exactly one frontmatter block
        self.assertEqual(self.entry()["filter_metadata"]["tags"], ["alpha"])


class RoundTripGuardTests(unittest.TestCase):
    def test_body_that_would_become_source_frontmatter_is_refused(self):
        body = "---\nlooks: like frontmatter\n---\n\ncontent\n"
        with self.assertRaises(ApplyRefused) as ctx:
            build_candidate(".md", b"no frontmatter here\n", body, hash_managed_body(body))
        self.assertEqual(ctx.exception.status, "not_round_trip_stable")

    def test_same_body_is_fine_below_existing_source_frontmatter(self):
        body = "---\nlooks: like frontmatter\n---\n\ncontent\n"
        candidate, _ = build_candidate(".md", b"---\nreal: fm\n---\n\nold\n", body, hash_managed_body(body))
        self.assertEqual(document_body(candidate.decode(), "md"), body)

    def test_txt_candidate_round_trips(self):
        body = "---\nnot: frontmatter in txt\n---\n"
        candidate, _ = build_candidate(".txt", b"old\n", body, hash_managed_body(body))
        self.assertEqual(document_body(candidate.decode(), "txt"), body)


if __name__ == "__main__":
    unittest.main()
