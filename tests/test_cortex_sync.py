"""Phase B: bidirectional sync, merge, rename, delete/tombstone, journal and
recovery. Everything runs against temp dirs, a directory vault and an
in-memory Qdrant."""
import asyncio
import io
import json
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from cortex import conflicts, identity, versions
from cortex.db import CortexError
from main import _decision_cli, apply_proposal_cli
from tests.cortex_fixtures import MD, CortexFixture

run = asyncio.run


def counts(db):
    return {t: db.one(f"SELECT COUNT(*) AS n FROM {t}")["n"] for t in
            ("documents", "revisions", "heads", "events", "conflicts", "tombstones")}


class BidirectionalTests(CortexFixture):
    def setUp(self):
        super().setUp()
        self.seed()

    def test_source_only_edit_propagates_to_vault_and_qdrant(self):
        self.edit_source("B\n", "B-new\n")
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("SOURCE_ONLY_CHANGED", "ok"))
        self.assertEqual(self.note_body(), "A\nB-new\nC\nD\nE\n")
        texts = [p.payload["text"] for p in self.points()]
        self.assertTrue(any("B-new" in t for t in texts))
        self.assertEqual({p.payload["path"] for p in self.points()}, {"local_ingest/plan.md"})
        self.assertEqual({p.payload.get("document_id") for p in self.points()}, {self.doc_id})
        self.assertEqual({p.payload.get("revision_id") for p in self.points()}, {self.doc()["current_revision_id"]})
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_vault_only_edit_goes_through_proposal_approve_apply(self):
        self.edit_note("E\n", "E-vault\n")
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("VAULT_ONLY_CHANGED", "proposal_pending"))
        self.assertNotIn("E-vault", self.source.read_text())  # nothing applied without approval
        pdir = self.root / "state" / "proposals"
        proposal_id = r["proposal_id"]
        with redirect_stdout(io.StringIO()):
            self.assertEqual(_decision_cli(pdir, r["proposal_id"], "approve", None, obsidian=self.vault,
                                           state_path=self.state, cortex_db=self.db.path), 0)
            self.assertEqual(apply_proposal_cli(r["proposal_id"], proposals_dir=pdir, state_path=self.state,
                                                obsidian=self.vault, cortex_db=self.db.path), 0)
        self.assertTrue(self.source.read_text().startswith("---\ntitle: Plan"))
        self.assertIn("E-vault", self.source.read_text())
        r = self.result()
        self.assertEqual(r["state"], "CONVERGED")
        rev = versions.get_revision(self.db, self.doc()["current_revision_id"])
        self.assertEqual(rev["origin"], "vault")
        self.assertEqual(rev["provenance"][0]["proposal_id"], proposal_id)
        self.assertEqual(rev["provenance"][0]["apply_actor"], "local")
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_non_overlapping_dual_edit_auto_merges_into_both_sides(self):
        self.edit_source("A\n", "A-src\n")
        self.edit_note("E\n", "E-vault\n")
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("BOTH_CHANGED", "ok"))
        expected = "A-src\nB\nC\nD\nE-vault\n"
        self.assertEqual(self.note_body(), expected)
        self.assertEqual(self.source.read_text(), "---\ntitle: Plan\nproject: kc\n---\n\n" + expected)
        self.assertTrue(all(Path(b).exists() for b in r["backups"]))
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_overlapping_dual_edit_is_a_conflict_and_writes_nothing(self):
        self.edit_source("C\n", "C-src\n")
        self.edit_note("C\n", "C-vault\n")
        src, note = self.source.read_text(), self.note_path.read_text()
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("CONFLICT", "conflict_recorded"))
        self.assertEqual((self.source.read_text(), self.note_path.read_text()), (src, note))
        c = conflicts.get_conflict(self.db, r["conflict_id"])
        self.assertEqual(c["regions"][0]["left"], "C-src\n")
        self.assertEqual(c["regions"][0]["right"], "C-vault\n")
        before = counts(self.db)
        again = self.result()
        self.assertEqual((again["status"], again["conflict_id"]), ("conflict_open", r["conflict_id"]))
        self.assertEqual(counts(self.db), before)  # no duplicate conflict/revisions

    def test_changed_files_supersede_an_open_conflict(self):
        self.edit_source("C\n", "C-src\n")
        self.edit_note("C\n", "C-vault\n")
        first = self.result()["conflict_id"]
        self.edit_note("C-vault\n", "C-vault2\n")
        second = self.result()
        self.assertNotEqual(second["conflict_id"], first)
        self.assertEqual(conflicts.get_conflict(self.db, first)["status"], "superseded")

    def test_repeat_sync_is_idempotent(self):
        before = counts(self.db)
        for _ in range(3):
            self.assertEqual(self.result()["state"], "IN_SYNC")
        self.assertEqual(counts(self.db), before)

    def test_dry_run_writes_nothing(self):
        self.edit_source("A\n", "A1\n")
        self.edit_note("E\n", "E1\n")
        files = {p: p.read_bytes() for p in self.root.rglob("*") if p.is_file() and "cortex.db" not in p.name}
        before = counts(self.db)
        report = self.sync(dry_run=True)
        self.assertEqual(report["documents"][0]["state"], "BOTH_CHANGED")
        self.assertIn("auto-merge", report["documents"][0]["planned_action"])
        self.assertEqual(counts(self.db), before)
        self.assertEqual({p: p.read_bytes() for p in files}, files)

    def test_human_edit_racing_a_write_is_never_overwritten(self):
        self.edit_source("A\n", "A1\n")
        real = self.engine._reindex

        def reindex_then_human_edits(doc):
            entry = real(doc)
            self.edit_note("E\n", "E-typed-meanwhile\n")
            return entry

        with patch.object(self.engine, "_reindex", side_effect=reindex_then_human_edits):
            r = self.result()
        self.assertEqual(r["status"], "stale_precondition")
        self.assertIn("E-typed-meanwhile", self.note_body())
        fail = [e for e in self.engine.journal.events(self.doc_id) if e["phase"] == "fail"]
        self.assertEqual(len(fail), 1)
        self.assertEqual(self.engine.journal.verify(), [])


class RenameTests(CortexFixture):
    def setUp(self):
        super().setUp()
        self.seed()

    def test_source_rename_keeps_identity_and_moves_note_and_index(self):
        old_note = self.doc()["vault_path"]
        new = self.src_dir / "sub" / "renamed.md"
        new.parent.mkdir()
        self.source.rename(new)
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("SOURCE_RENAMED", "ok"))
        doc = self.doc()
        self.assertEqual(doc["document_id"], self.doc_id)
        self.assertEqual(doc["source_path"], str(new))
        self.assertNotEqual(doc["vault_path"], old_note)
        self.assertTrue(doc["vault_path"].startswith("Knowledge Cortex/Managed/renamed-"))
        self.assertFalse((self.root / "vault" / old_note).exists())
        content = (self.root / "vault" / doc["vault_path"]).read_text()
        self.assertIn(f"cortex_source_path: {new}", content)
        self.assertIn(f"cortex_document_id: {self.doc_id}", content)
        self.assertEqual({p.payload["path"] for p in self.points()}, {"local_ingest/renamed.md"})
        state = json.loads(self.state.read_text())["local"]
        self.assertIn(str(new), state)
        self.assertNotIn(str(self.source.resolve()), state)
        self.source = new
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_vault_rename_updates_relationship_only(self):
        moved = self.root / "vault" / "Knowledge Cortex" / "Managed" / "Projects" / "plan-renamed.md"
        moved.parent.mkdir()
        self.note_path.rename(moved)
        src_before = self.source.read_bytes()
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("VAULT_RENAMED", "ok"))
        self.assertEqual(self.doc()["vault_path"], "Knowledge Cortex/Managed/Projects/plan-renamed.md")
        self.assertEqual(self.source.read_bytes(), src_before)
        self.assertEqual(self.doc()["document_id"], self.doc_id)
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_rename_collision_moves_nothing(self):
        new = self.src_dir / "other.md"
        target = self.root / "vault" / "Knowledge Cortex" / "Managed" / \
            f"other-{__import__('ingest.vault_writer', fromlist=['x']).source_id_for(str(new))}.md"
        target.write_text("a human note already here\n")
        old_note = self.note_path
        self.source.rename(new)
        r = self.result()
        self.assertEqual(r["status"], "collision")
        self.assertTrue(old_note.exists())
        self.assertEqual(target.read_text(), "a human note already here\n")
        self.assertEqual(self.doc()["source_path"], str(self.source.resolve()))

    def test_ambiguous_source_rename_is_not_guessed(self):
        copy_a, copy_b = self.src_dir / "a.md", self.src_dir / "b.md"
        copy_a.write_bytes(self.source.read_bytes())
        copy_b.write_bytes(self.source.read_bytes())
        self.source.unlink()
        r = self.result()
        self.assertEqual(r["state"], "INVALID")
        self.assertEqual(len(r["candidates"]), 2)

    def test_unmanaged_moved_note_is_not_adopted(self):
        content = self.note_path.read_text().replace("cortex_managed: true", "cortex_managed: false")
        elsewhere = self.root / "vault" / "Knowledge Cortex" / "Managed" / "x.md"
        elsewhere.write_text(content)
        self.note_path.unlink()
        r = self.result()
        self.assertEqual(r["state"], "VAULT_DELETED")


class DeleteTests(CortexFixture):
    def setUp(self):
        super().setUp()
        self.seed()

    def test_source_deletion_tombstone_propagation_and_restore(self):
        self.source.unlink()
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("SOURCE_DELETED", "tombstoned"))
        self.assertTrue(self.note_path.exists())  # nothing deleted yet
        tomb = self.engine.get_tombstone(r["tombstone_id"])
        self.assertEqual((tomb["deleted_from"], tomb["deleted_by"], tomb["status"]), ("source", "local", "detected"))
        note_path = self.note_path
        result = run(self.engine.approve_tombstone("local", r["tombstone_id"]))
        self.assertEqual(result["status"], "propagated")
        self.assertFalse(note_path.exists())
        self.assertTrue((self.root / "vault" / result["trash"]["vault"]).exists())
        self.assertEqual(self.doc()["status"], "tombstoned")
        self.assertEqual(self.points(), [])
        self.assertEqual(self.result()["state"], "TOMBSTONED")
        restored = run(self.engine.restore_tombstone("local", r["tombstone_id"]))
        self.assertEqual(restored["status"], "restored")
        self.assertEqual(self.source.read_text(), MD)  # frontmatter + body reconstructed
        self.assertTrue(note_path.exists())
        self.assertEqual(self.result()["state"], "IN_SYNC")
        self.assertTrue(self.points())

    def test_vault_deletion_moves_source_to_recovery_area_and_restores(self):
        self.note_path.unlink()
        r = self.result()
        self.assertEqual(r["state"], "VAULT_DELETED")
        result = run(self.engine.approve_tombstone("local", r["tombstone_id"]))
        trash = Path(result["trash"]["source"])
        self.assertFalse(self.source.exists())
        self.assertEqual(trash.read_text(), MD)
        self.assertTrue(str(trash).startswith(str(self.root / "state" / "trash")))
        run(self.engine.restore_tombstone("local", r["tombstone_id"]))
        self.assertEqual(self.source.read_text(), MD)
        self.assertFalse(trash.exists())
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_restore_before_propagation_recreates_the_note(self):
        self.note_path.unlink()
        tid = self.result()["tombstone_id"]
        run(self.engine.restore_tombstone("local", tid))
        self.assertEqual(self.note_body(), "A\nB\nC\nD\nE\n")
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_purge_is_admin_only_and_only_after_propagation(self):
        from cortex import users
        users.add_user(self.db, "local", "bob", "member")
        self.note_path.unlink()
        tid = self.result()["tombstone_id"]
        with self.assertRaises(CortexError):
            run(self.engine.purge_tombstone("local", tid))  # not propagated yet
        result = run(self.engine.approve_tombstone("local", tid))
        with self.assertRaises(CortexError):
            run(self.engine.purge_tombstone("bob", tid))
        run(self.engine.purge_tombstone("local", tid))
        self.assertFalse(Path(result["trash"]["source"]).exists())
        self.assertEqual(self.engine.get_tombstone(tid)["status"], "purged")

    def test_approval_refused_if_the_deleted_side_came_back(self):
        self.note_path.unlink()
        tid = self.result()["tombstone_id"]
        run(self.engine.restore_tombstone("local", tid))
        with self.assertRaises(CortexError):
            run(self.engine.approve_tombstone("local", tid))
        self.assertTrue(self.source.exists())

    def test_backend_without_moves_fails_closed(self):
        self.source.unlink()
        tid = self.result()["tombstone_id"]
        self.engine.vault.supports_moves = False
        try:
            with self.assertRaises(CortexError) as ctx:
                run(self.engine.approve_tombstone("local", tid))
        finally:
            self.engine.vault.supports_moves = True
        self.assertEqual(ctx.exception.code, "backend_unsupported")
        self.assertTrue(self.note_path.exists())


class UnsupportedFormatTests(CortexFixture):
    def test_docx_source_vault_edit_is_preserved_never_reverse_written(self):
        from docx import Document
        d = Document()
        d.add_paragraph("First line")
        d.add_paragraph("Second line")
        path = self.src_dir / "report.docx"
        d.save(path)
        with redirect_stdout(io.StringIO()):
            from main import index_to_qdrant
            index_to_qdrant(str(path), None, self.state, store=self.store, write_vault=True, obsidian=self.vault)
        self.source = path
        self.sync()
        self.doc_id = self.db.one("SELECT document_id FROM documents")["document_id"]
        self.assertEqual(self.doc()["reverse_writable"], 0)
        raw = path.read_bytes()
        self.edit_note("- Second line", "- Second line edited in vault")
        r = self.result()
        self.assertEqual((r["state"], r["status"]), ("VAULT_ONLY_CHANGED", "preserved"))
        self.assertEqual(path.read_bytes(), raw)
        self.assertIn("edited in vault", self.note_body())
        with self.assertRaises(CortexError) as ctx:
            run(conflicts.edit_document(self.engine, "local", self.doc_id, "x\n", self.doc()["current_revision_id"]))
        self.assertIn(ctx.exception.code, ("unsupported_source_type", "not_in_sync"))


class RecoveryTests(CortexFixture):
    def setUp(self):
        super().setUp()
        self.seed()

    def test_crash_after_file_writes_before_commit_is_recovered(self):
        self.edit_source("A\n", "A1\n")
        real_finish = self.engine.journal.finish

        def crash_on_commit(op_id, phase, actor, data, db_changes=None):
            if phase == "commit":
                raise KeyboardInterrupt("simulated crash")
            return real_finish(op_id, phase, actor, data, db_changes)

        with patch.object(self.engine.journal, "finish", side_effect=crash_on_commit), \
                patch.object(self.engine, "_compensate", side_effect=lambda *a: []):
            with self.assertRaises(KeyboardInterrupt):
                self.sync()
        self.assertEqual(self.note_body(), "A1\nB\nC\nD\nE\n")  # file written
        self.assertEqual(len(self.engine.journal.pending()), 1)
        old_rev = self.doc()["current_revision_id"]
        outcomes = run(self.engine.recover("local"))
        self.assertEqual(outcomes[0]["result"], "recovered")
        self.assertNotEqual(self.doc()["current_revision_id"], old_rev)
        self.assertEqual(self.engine.journal.pending(), [])
        self.assertEqual(self.result()["state"], "IN_SYNC")
        self.assertEqual(versions.verify_graph(self.db), [])
        self.assertEqual(self.engine.journal.verify(), [])

    def test_crash_before_any_write_rolls_back(self):
        self.edit_source("A\n", "A1\n")
        with patch.object(self.engine, "_vault_write", side_effect=KeyboardInterrupt), \
                patch.object(self.engine.journal, "finish", side_effect=KeyboardInterrupt):
            with self.assertRaises(KeyboardInterrupt):
                self.sync()
        outcomes = run(self.engine.recover("local"))
        self.assertEqual(outcomes[0]["result"], "rolled_back")
        r = self.result()  # the next normal sync does the work
        self.assertEqual((r["state"], r["status"]), ("SOURCE_ONLY_CHANGED", "ok"))
        self.assertEqual(self.result()["state"], "IN_SYNC")

    def test_partial_multi_file_write_is_compensated(self):
        self.edit_source("A\n", "A1\n")
        self.edit_note("E\n", "E1\n")
        note_before, src_before = self.note_path.read_text(), self.source.read_text()
        real_vault_write = self.engine._vault_write

        async def vault_write_then_fail(*args, **kw):
            await real_vault_write(*args, **kw)  # the note IS rewritten...
            raise OSError("disk full")           # ...then the operation fails

        with patch.object(self.engine, "_vault_write", side_effect=vault_write_then_fail):
            r = self.result()
        self.assertNotEqual(r["status"], "ok")
        self.assertEqual(self.source.read_text(), src_before)
        self.assertEqual(self.note_path.read_text(), note_before)
        fail = [e for e in self.engine.journal.events(self.doc_id) if e["phase"] == "fail"][-1]
        self.assertIn("disk full", fail["data"]["error"])
        self.assertTrue(fail["data"]["compensation"])  # the rewritten files were put back

    def test_interrupted_merge_after_source_write_is_rolled_back_from_backups(self):
        self.edit_source("A\n", "A1\n")
        self.edit_note("E\n", "E1\n")
        src_before, note_before = self.source.read_bytes(), self.note_path.read_text()
        real_finish = self.engine.journal.finish

        def crash(op_id, phase, actor, data, db_changes=None):
            if phase in ("commit", "fail"):
                raise KeyboardInterrupt
            return real_finish(op_id, phase, actor, data, db_changes)

        with patch.object(self.engine, "_vault_write", side_effect=KeyboardInterrupt), \
                patch.object(self.engine.journal, "finish", side_effect=crash), \
                patch.object(self.engine, "_compensate", side_effect=lambda *a: []):
            with self.assertRaises(KeyboardInterrupt):
                self.sync()
        self.assertNotEqual(self.source.read_bytes(), src_before)  # source was replaced, vault not
        outcomes = run(self.engine.recover("local"))
        self.assertEqual(outcomes[0]["result"], "rolled_back_from_backups")
        self.assertEqual(self.source.read_bytes(), src_before)
        self.assertEqual(self.note_path.read_text(), note_before)
        self.assertEqual(self.result()["state"], "BOTH_CHANGED")
